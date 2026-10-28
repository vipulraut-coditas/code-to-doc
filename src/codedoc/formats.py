"""Output formats: Markdown, Word (.docx) and Google Docs.

Every document is rendered to Markdown first (``render.py``); the other formats
are converted from that Markdown, so choosing more formats costs no extra LLM calls.

    md    -> written as-is
    docx  -> Markdown parsed with markdown-it, rebuilt with python-docx
    gdoc  -> the .docx is uploaded to Google Drive, which converts it to a Google Doc
"""

from __future__ import annotations

import io
import json
import os
import re
import uuid
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import IO, Any

from codedoc.config import config_dir


@dataclass(frozen=True)
class OutputFormat:
    id: str
    label: str
    description: str
    extension: str = ""   # empty = not a local file
    local: bool = True


FORMATS: list[OutputFormat] = [
    OutputFormat("md", "Markdown (.md)", "plain text, great for Git repos and wikis", ".md"),
    OutputFormat("docx", "Word document (.docx)", "open and edit in Microsoft Word / LibreOffice", ".docx"),
    OutputFormat("gdoc", "Google Docs", "uploaded to your Google Drive (Google sign-in needed)", local=False),
]

FORMAT_IDS = [f.id for f in FORMATS]
_BY_ID = {f.id: f for f in FORMATS}


def get_format(format_id: str) -> OutputFormat:
    try:
        return _BY_ID[format_id]
    except KeyError:
        raise ValueError(f"Unknown format '{format_id}'. Choose from: {', '.join(FORMAT_IDS)}") from None


def parse_formats(value: str | list[str] | tuple[str, ...]) -> list[str]:
    """'md,docx' / ['docx'] / 'all' -> validated ids in menu order, without duplicates."""
    items = value.split(",") if isinstance(value, str) else list(value)
    items = [i.strip().lower() for i in items if i.strip()]
    items = ["docx" if i == "word" else "gdoc" if i in {"google", "gdocs"} else i for i in items]
    if "all" in items:
        return list(FORMAT_IDS)
    if not items:
        raise ValueError(f"Choose at least one format: {', '.join(FORMAT_IDS)}")
    for i in items:
        get_format(i)
    return [f for f in FORMAT_IDS if f in items]


def needs_output_dir(formats: list[str]) -> bool:
    return any(get_format(f).local for f in formats)


# ---------------------------------------------------------------- Markdown -> Word
# characters XML 1.0 (and so python-docx) cannot store; a model reply containing one would lose the .docx
_XML_INVALID = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def markdown_to_docx(markdown: str, target: str | Path | IO[bytes], title: str = "") -> None:
    """Convert the Markdown our renderers produce (headings, lists, tables, code, quotes) to .docx."""
    from docx import Document
    from docx.shared import Pt
    from markdown_it import MarkdownIt

    markdown, title = _XML_INVALID.sub("", markdown), _XML_INVALID.sub("", title)
    doc = Document()
    if title:
        doc.core_properties.title = title
    tokens = MarkdownIt("commonmark").enable("table").parse(markdown)
    lists: list[dict[str, Any]] = []   # stack of {"ordered": bool, "n": int}
    in_quote = False
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        kind = tok.type
        if kind == "heading_open":
            level = int(tok.tag[1]) - 1   # "#" -> Word "Title", "##" -> "Heading 1", ...
            _add_inline(doc.add_heading(level=level), tokens[i + 1])
            i += 3
            continue
        if kind in ("bullet_list_open", "ordered_list_open"):
            lists.append({"ordered": kind == "ordered_list_open", "n": int(tok.attrs.get("start", 1) or 1) - 1})
        elif kind in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif kind == "list_item_open" and lists:
            lists[-1]["n"] += 1
        elif kind == "blockquote_open":
            in_quote = True
        elif kind == "blockquote_close":
            in_quote = False
        elif kind == "paragraph_open":
            inline = tokens[i + 1]
            if lists:
                current = lists[-1]
                para = doc.add_paragraph(style="List Paragraph")
                para.paragraph_format.left_indent = Pt(18 * len(lists))
                para.paragraph_format.first_line_indent = Pt(-14)
                para.add_run(f"{current['n']}. " if current["ordered"] else "• ")
            else:
                para = doc.add_paragraph(style="Quote" if in_quote else None)
            _add_inline(para, inline)
            i += 3
            continue
        elif kind in ("fence", "code_block"):
            for line in tok.content.rstrip("\n").split("\n"):
                run = doc.add_paragraph().add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
        elif kind == "table_open":
            i = _add_table(doc, tokens, i)
            continue
        elif kind == "hr":
            doc.add_paragraph()
        i += 1
    doc.save(target)


def _add_inline(paragraph, inline_token) -> None:
    from docx.shared import Pt

    bold = italic = False
    link: str | None = None
    for child in inline_token.children or []:
        kind = child.type
        if kind == "strong_open":
            bold = True
        elif kind == "strong_close":
            bold = False
        elif kind == "em_open":
            italic = True
        elif kind == "em_close":
            italic = False
        elif kind == "link_open":
            link = child.attrs.get("href")
        elif kind == "link_close":
            if link:
                paragraph.add_run(f" ({link})")
            link = None
        elif kind in ("softbreak", "hardbreak"):
            paragraph.add_run().add_break()
        elif kind in ("text", "code_inline", "html_inline"):
            run = paragraph.add_run(child.content)
            run.bold, run.italic = bold or None, italic or None
            if kind == "code_inline":
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)


def _add_table(doc, tokens, i: int) -> int:
    """Collect the rows of a Markdown table starting at ``tokens[i]`` and add a Word table. Returns the next index."""
    rows: list[tuple[bool, list]] = []   # (is_header, [inline tokens])
    header = False
    while tokens[i].type != "table_close":
        kind = tokens[i].type
        if kind == "thead_open":
            header = True
        elif kind == "thead_close":
            header = False
        elif kind == "tr_open":
            rows.append((header, []))
        elif kind == "inline":
            rows[-1][1].append(tokens[i])
        i += 1
    if rows:
        cols = max(len(cells) for _, cells in rows)
        table = doc.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"
        for r, (is_header, cells) in enumerate(rows):
            for c, inline in enumerate(cells):
                para = table.cell(r, c).paragraphs[0]
                _add_inline(para, inline)
                if is_header:
                    for run in para.runs:
                        run.bold = True
        doc.add_paragraph()
    return i + 1


# ---------------------------------------------------------------- Google Docs
GOOGLE_SCOPES = ["https://www.googleapis.com/auth/drive.file"]  # only files codedoc creates
GOOGLE_PACKAGE = "google-auth-oauthlib"
GOOGLE_SETUP_HELP = """To upload to Google Docs, codedoc needs a Google OAuth client (one-time setup):
  1. Open https://console.cloud.google.com/ and create (or pick) a project.
  2. APIs & Services > Library: enable the "Google Drive API".
  3. APIs & Services > OAuth consent screen: choose External and add your Google account as a test user.
  4. APIs & Services > Credentials > Create credentials > OAuth client ID > "Desktop app", then download the JSON.
  5. Save it as {path}
     (or point the CODEDOC_GOOGLE_CREDENTIALS environment variable at it)."""


class GoogleSetupError(RuntimeError):
    """Google Docs output was chosen but the libraries or the OAuth client file are missing."""


def google_credentials_file() -> Path:
    return Path(os.environ.get("CODEDOC_GOOGLE_CREDENTIALS") or config_dir() / "google_credentials.json")


def google_token_file() -> Path:
    return config_dir() / "google_token.json"


def google_libraries_installed() -> bool:
    import importlib.util

    return importlib.util.find_spec("google_auth_oauthlib") is not None


def google_authorize():
    """Return Google credentials: saved token, refreshed token, or a one-time browser sign-in."""
    if not google_libraries_installed():
        raise GoogleSetupError(f"Google Docs output needs the '{GOOGLE_PACKAGE}' package: "
                               f"pip install \"codedoc-agent[gdocs]\"")
    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow

    token = google_token_file()
    creds = None
    if token.exists():
        try:
            creds = Credentials.from_authorized_user_file(str(token), GOOGLE_SCOPES)
        except (ValueError, json.JSONDecodeError):
            creds = None
    if creds and creds.valid:
        return creds
    if creds and creds.expired and creds.refresh_token:
        try:
            creds.refresh(Request())
        except Exception:  # revoked or expired refresh token: sign in again
            creds = None
    if not creds or not creds.valid:
        client = google_credentials_file()
        if not client.exists():
            raise GoogleSetupError(GOOGLE_SETUP_HELP.format(path=client))
        flow = InstalledAppFlow.from_client_secrets_file(str(client), GOOGLE_SCOPES)
        creds = flow.run_local_server(port=0, open_browser=True,
                                      authorization_prompt_message="Opening your browser to sign in to Google: {url}",
                                      success_message="codedoc is signed in to Google. You can close this tab.")
    token.parent.mkdir(parents=True, exist_ok=True)
    token.write_text(creds.to_json(), encoding="utf-8")
    try:
        os.chmod(token, 0o600)  # refresh token: owner only (best effort; no-op on Windows)
    except OSError:
        pass
    return creds


def upload_to_google_docs(creds, title: str, docx_bytes: bytes, folder_id: str | None = None) -> str:
    """Upload a .docx to Drive, converted to a native Google Doc. Returns the document's web link."""
    from google.auth.transport.requests import AuthorizedSession

    metadata: dict[str, Any] = {"name": title, "mimeType": "application/vnd.google-apps.document"}
    if folder_id:
        metadata["parents"] = [folder_id]
    boundary = uuid.uuid4().hex
    body = (f"--{boundary}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n"
            f"{json.dumps(metadata)}\r\n--{boundary}\r\n"
            "Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n"
            ).encode() + docx_bytes + f"\r\n--{boundary}--\r\n".encode()
    response = AuthorizedSession(creds).post(
        "https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart&fields=id,webViewLink",
        data=body, headers={"Content-Type": f"multipart/related; boundary={boundary}"}, timeout=120)
    if response.status_code >= 400:
        raise RuntimeError(f"Google Drive upload failed ({response.status_code}): {response.text[:300]}")
    data = response.json()
    return data.get("webViewLink") or f"https://docs.google.com/document/d/{data['id']}/edit"


# ---------------------------------------------------------------- export one document
def _backup_existing(path: Path) -> str | None:
    """Rename an existing file to ``<stem>.backup-<timestamp><ext>`` so it is never silently lost."""
    if not path.exists():
        return None
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = path.with_name(f"{path.stem}.backup-{stamp}{path.suffix}")
    n = 1
    while backup.exists():
        n += 1
        backup = path.with_name(f"{path.stem}.backup-{stamp}-{n}{path.suffix}")
    path.replace(backup)
    return f"Existing {path.name} was kept as {backup.name}"


def export_document(markdown: str, *, formats: list[str], output_dir: Path | None, filename: str,
                    title: str, google_creds=None, google_folder: str | None = None) -> tuple[dict[str, str], list[str]]:
    """Write/upload one document in every chosen format.

    Returns ``({format_id: path or URL}, warnings)``; a failing format does not stop the others.
    """
    outputs: dict[str, str] = {}
    warnings: list[str] = []
    stem = Path(filename).stem
    for fmt_id in formats:
        fmt = get_format(fmt_id)
        try:
            if fmt.id == "md":
                path = output_dir / f"{stem}.md"
                if note := _backup_existing(path):
                    warnings.append(note)
                path.write_text(markdown, encoding="utf-8")
                outputs[fmt.id] = str(path)
            elif fmt.id == "docx":
                path = output_dir / f"{stem}.docx"
                if note := _backup_existing(path):
                    warnings.append(note)
                markdown_to_docx(markdown, path, title=title)
                outputs[fmt.id] = str(path)
            elif fmt.id == "gdoc":
                buffer = io.BytesIO()
                markdown_to_docx(markdown, buffer, title=title)
                outputs[fmt.id] = upload_to_google_docs(google_creds, title, buffer.getvalue(), google_folder)
        except Exception as exc:
            warnings.append(f"Could not create the {fmt.label} version of {title}: {str(exc)[:200]}")
    return outputs, warnings
