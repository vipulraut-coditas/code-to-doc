import json

import pytest
from docx import Document

from codedoc import formats
from codedoc.api import generate_docs
from codedoc.cli import build_parser
from codedoc.config import Settings
from codedoc.formats import FORMAT_IDS, GoogleSetupError, markdown_to_docx, parse_formats

from conftest import FakeBA


def settings():
    return Settings(backend="api", provider="anthropic", fast_model="f", smart_model="s", concurrency=3)


@pytest.fixture
def fake_google(monkeypatch):
    """Replace Google sign-in and upload; records every upload."""
    uploads = []

    def upload(creds, title, docx_bytes, folder_id=None):
        uploads.append({"creds": creds, "title": title, "bytes": docx_bytes, "folder": folder_id})
        return f"https://docs.google.com/document/d/doc{len(uploads)}/edit"

    monkeypatch.setattr(formats, "google_authorize", lambda: "CREDS")
    monkeypatch.setattr(formats, "upload_to_google_docs", upload)
    return uploads


def test_format_parsing():
    assert parse_formats("docx, md") == ["md", "docx"]
    assert parse_formats("word,google") == ["docx", "gdoc"]
    assert parse_formats("all") == FORMAT_IDS
    with pytest.raises(ValueError, match="Choose from"):
        parse_formats("pdf")
    assert build_parser().parse_args(["generate", "."]).format == ["md"]
    assert build_parser().parse_args(["generate", ".", "-f", "docx,gdoc"]).format == ["docx", "gdoc"]
    with pytest.raises(SystemExit):
        build_parser().parse_args(["generate", ".", "--format", "pdf"])


def test_markdown_to_docx_keeps_structure(tmp_path):
    md = ("# Title\n## Project\n\n> A **draft** note.\n\n## 1. Intro\nSome *text* with `code`.\n\n"
          "- first\n- second\n\n1. step one\n2. step two\n\n"
          "| ID | Requirement |\n|---|---|\n| FR-001 | Users \\| admins can log in |\n\n"
          "```mermaid\nflowchart LR\n  A --> B\n```\n")
    out = tmp_path / "x.docx"
    markdown_to_docx(md, out, title="Doc")

    doc = Document(out)
    paras = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert ("Title", "Title") in paras and ("Heading 1", "1. Intro") in paras
    assert ("Quote", "A draft note.") in paras
    assert ("List Paragraph", "• first") in paras and ("List Paragraph", "2. step two") in paras
    assert any(t == "flowchart LR" for _, t in paras)
    bold = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "draft" in bold
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["ID", "Requirement"]
    assert table.rows[1].cells[1].text == "Users | admins can log in"  # escaped pipe restored
    assert doc.core_properties.title == "Doc"


def test_markdown_and_word_in_one_run(sample_repo, tmp_path, fake_llm):
    result = generate_docs(str(sample_repo), ["srs", "api"], tmp_path / "out", settings(), formats="md,docx",
                           fast_llm=fake_llm, smart_llm=fake_llm)

    names = sorted(p.name for p in (tmp_path / "out").iterdir())
    assert names == ["API.docx", "API.md", "SRS.docx", "SRS.md"]
    srs = result.get("srs")
    assert set(srs.outputs) == {"md", "docx"} and srs.output_path.name == "SRS.md"
    heading = Document(tmp_path / "out" / "SRS.docx").paragraphs[0].text
    assert heading == "Software Requirements Specification (SRS)"
    assert fake_llm.calls["write_srs"] == 1  # extra formats cost no LLM calls


def test_google_docs_upload(sample_repo, tmp_path, fake_llm, fake_google):
    result = generate_docs(str(sample_repo), ["srs", "user_manual"], tmp_path / "out", settings(), formats=["gdoc"],
                           google_folder="FOLDER1", fast_llm=fake_llm, smart_llm=fake_llm)

    assert [u["title"] for u in fake_google] == ["shop - Software Requirements Specification (SRS)",
                                                 "shop - User Manual"]
    assert all(u["creds"] == "CREDS" and u["folder"] == "FOLDER1" for u in fake_google)
    assert fake_google[0]["bytes"][:2] == b"PK"  # a real .docx (zip) is uploaded
    assert result.get("srs").outputs == {"gdoc": "https://docs.google.com/document/d/doc1/edit"}
    assert result.output_dir is None and not (tmp_path / "out").exists()  # Google only: no local files


def test_failed_upload_keeps_local_files(sample_repo, tmp_path, fake_llm, monkeypatch):
    monkeypatch.setattr(formats, "google_authorize", lambda: "CREDS")

    def broken(*a, **k):
        raise RuntimeError("quota exceeded")

    monkeypatch.setattr(formats, "upload_to_google_docs", broken)
    result = generate_docs(str(sample_repo), ["srs"], tmp_path / "out", settings(), formats="md,gdoc",
                           fast_llm=fake_llm, smart_llm=fake_llm)

    assert (tmp_path / "out" / "SRS.md").exists()
    assert set(result.get("srs").outputs) == {"md"}
    assert any("Google Docs" in w and "quota exceeded" in w for w in result.warnings)


def test_google_setup_problem_stops_before_any_llm_call(sample_repo, tmp_path):
    llm = FakeBA(calls={})
    with pytest.raises(GoogleSetupError, match="Google Drive API"):  # no OAuth client file in CODEDOC_HOME
        generate_docs(str(sample_repo), ["srs"], tmp_path / "out", settings(), formats="gdoc",
                      fast_llm=llm, smart_llm=llm)
    assert llm.calls == {}


def test_upload_request_converts_to_google_doc(monkeypatch):
    sent = {}

    class FakeResponse:
        status_code = 200

        def json(self):
            return {"id": "abc", "webViewLink": "https://docs.google.com/document/d/abc/edit"}

    class FakeSession:
        def __init__(self, creds):
            sent["creds"] = creds

        def post(self, url, data, headers, timeout):
            sent.update(url=url, data=data, headers=headers)
            return FakeResponse()

    import google.auth.transport.requests as transport
    monkeypatch.setattr(transport, "AuthorizedSession", FakeSession)

    link = formats.upload_to_google_docs("CREDS", "My Doc", b"PK-docx-bytes", folder_id="F1")

    assert link == "https://docs.google.com/document/d/abc/edit"
    assert "uploadType=multipart" in sent["url"]
    boundary = sent["headers"]["Content-Type"].split("boundary=")[1]
    metadata = json.loads(sent["data"].split(f"--{boundary}".encode())[1].split(b"\r\n\r\n", 1)[1])
    assert metadata == {"name": "My Doc", "mimeType": "application/vnd.google-apps.document", "parents": ["F1"]}
    assert b"PK-docx-bytes" in sent["data"]
